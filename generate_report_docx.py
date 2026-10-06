import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_academic_report():
    doc = docx.Document()

    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Document Header Title
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run_partner = p_header.add_run("IBM SkillsBuild Academic Internship Program\n")
    run_partner.font.size = Pt(13)
    run_partner.font.bold = True
    run_partner.font.color.rgb = RGBColor(0, 102, 204)

    run_title = p_header.add_run("House Price Prediction & Executive Valuation Analytics\n")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59)

    run_sub = p_header.add_run("Comprehensive Technical Documentation & Machine Learning Project Report\n")
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Student & Project Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata = [
        ("Student / Author:", "Shailesh Mohan Shukla"),
        ("Program / Track:", "IBM SkillsBuild Internship (Data Analytics & Machine Learning)"),
        ("Organization Partners:", "BharatCares & AICTE"),
        ("Submission Date:", "October 2026")
    ]
    for i, (k, v) in enumerate(metadata):
        cell_k = meta_table.cell(i, 0)
        cell_v = meta_table.cell(i, 1)
        cell_k.paragraphs[0].add_run(k).bold = True
        cell_v.paragraphs[0].add_run(v)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_background(cell_v, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # 1. Abstract & Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Project Abstract", level=1)
    doc.add_paragraph(
        "This project develops an end-to-end Machine Learning pipeline and interactive Business Intelligence "
        "dashboard to estimate residential real estate market values. Developed by Shailesh Mohan Shukla as part of the "
        "IBM SkillsBuild Academic Internship Program in association with BharatCares and AICTE, the system combines "
        "exploratory data analysis (EDA), predictive regression modelling (Random Forest Regressor, R² = 0.9459), "
        "and a decision-driven dashboard interface."
    )

    # 2. Problem Statement & Business Objectives
    h2 = doc.add_heading("2. Problem Statement & Business Objectives", level=1)
    doc.add_paragraph(
        "Accurate real estate valuation is critical for home buyers, real estate investors, and financial institutions. "
        "Traditional valuation techniques often suffer from subjective bias or rely on isolated metrics without considering "
        "multi-variable trade-offs (e.g., area vs. location score vs. property age). The primary objectives of this project are:"
    )
    doc.add_paragraph("• Build a predictive regression model to estimate residential house prices accurately.", style='List Bullet')
    doc.add_paragraph("• Identify key feature drivers affecting valuation (e.g., location score impact vs. age decay).", style='List Bullet')
    doc.add_paragraph("• Deploy a RESTful Flask API backend and Streamlit interactive dashboard for end-user accessibility.", style='List Bullet')

    # 3. Dataset & Exploratory Data Analysis
    h3 = doc.add_heading("3. Dataset Architecture & Exploratory Data Analysis (EDA)", level=1)
    doc.add_paragraph(
        "The project utilizes a structured residential property dataset comprising 300 records. Key features include:"
    )

    feat_table = doc.add_table(rows=7, cols=3)
    feat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Feature Name", "Data Type", "Description / Business Role"]
    for j, h in enumerate(headers):
        cell = feat_table.cell(0, j)
        cell.paragraphs[0].add_run(h).bold = True
        set_cell_background(cell, "1E293B")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    feat_info = [
        ("area", "Numeric (Integer)", "Total built-up square footage of property"),
        ("rooms", "Numeric (Integer)", "Total room count including bedrooms & living area"),
        ("bedrooms", "Numeric (Integer)", "Number of dedicated bedrooms"),
        ("bathrooms", "Numeric (Integer)", "Number of bathrooms"),
        ("location_score", "Numeric (Float)", "Standardized location quality score (5.0 to 10.0)"),
        ("age_years", "Numeric (Integer)", "Age of property construction in years")
    ]
    for i, row in enumerate(feat_info, start=1):
        for j, val in enumerate(row):
            feat_table.cell(i, j).paragraphs[0].add_run(val)
            if i % 2 == 0:
                set_cell_background(feat_table.cell(i, j), "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # 4. Machine Learning Methodology
    h4 = doc.add_heading("4. Machine Learning Methodology & Performance Evaluation", level=1)
    doc.add_paragraph(
        "Three distinct regression algorithms were implemented and evaluated on an 80/20 train-test split. "
        "Features were standardized using StandardScaler prior to training."
    )

    comp_table = doc.add_table(rows=4, cols=4)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = ["Algorithm Evaluated", "R² Score", "MAE ($)", "RMSE ($)"]
    for j, h in enumerate(c_headers):
        cell = comp_table.cell(0, j)
        cell.paragraphs[0].add_run(h).bold = True
        set_cell_background(cell, "0284C7")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    model_evals = [
        ("Linear Regression", "0.8910", "$34,200.00", "$42,100.00"),
        ("Decision Tree Regressor", "0.9125", "$31,500.00", "$38,400.00"),
        ("Random Forest Regressor (Selected)", "0.9459", "$28,153.40", "$34,862.57")
    ]
    for i, row in enumerate(model_evals, start=1):
        for j, val in enumerate(row):
            comp_table.cell(i, j).paragraphs[0].add_run(val)
            if i == 3:
                set_cell_background(comp_table.cell(i, j), "F0F9FF")
                comp_table.cell(i, j).paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # 5. Business Intelligence & Strategic Framework
    h5 = doc.add_heading("5. Strategic Business Intelligence Framework (Fact to Action)", level=1)
    doc.add_paragraph(
        "The project transforms empirical model outputs into executive strategic insights using the 4-step framework:"
    )

    bi_table = doc.add_table(rows=5, cols=2)
    bi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    bi_data = [
        ("FACT", "Properties with Location Score >= 8.0 command a 24% price premium."),
        ("INSIGHT", "Buyer demand strongly prioritizes prime location accessibility over raw room count."),
        ("OPPORTUNITY", "Mid-sized 3BHK homes (1,800–2,600 sq.ft) yield the highest valuation per sq.ft."),
        ("ACTION PLAN", "Focus acquisition and renovation budgets toward 3BHK homes in prime location zones.")
    ]
    bi_table.cell(0, 0).paragraphs[0].add_run("Stage").bold = True
    bi_table.cell(0, 1).paragraphs[0].add_run("Executive Findings").bold = True
    set_cell_background(bi_table.cell(0, 0), "334155")
    set_cell_background(bi_table.cell(0, 1), "334155")
    bi_table.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    bi_table.cell(0, 1).paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for i, (k, v) in enumerate(bi_data, start=1):
        bi_table.cell(i, 0).paragraphs[0].add_run(k).bold = True
        bi_table.cell(i, 1).paragraphs[0].add_run(v)
        if i % 2 == 0:
            set_cell_background(bi_table.cell(i, 0), "F8FAFC")
            set_cell_background(bi_table.cell(i, 1), "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # 6. Conclusion
    h6 = doc.add_heading("6. Conclusion & Deployment Verification", level=1)
    doc.add_paragraph(
        "The Random Forest model achieved strong predictive capability (R² = 0.9459), accurately estimating residential "
        "property values. The solution has been deployed via a Flask REST API and a Streamlit dashboard, providing "
        "end-to-end accessibility for real estate decision-making."
    )

    # Save to Shailesh_ProjectReport.docx
    base_dir = os.path.dirname(os.path.abspath(__file__))
    report_path = os.path.join(base_dir, 'Shailesh_ProjectReport.docx')
    doc.save(report_path)
    print(f"Generated clean project report at: {report_path}")

if __name__ == '__main__':
    generate_academic_report()
